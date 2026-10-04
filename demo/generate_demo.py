"""Regenerate the three synthetic inputs. Run from any working directory."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
HEADERS = [
    "asset_id", "device_type", "vendor", "model", "status", "location",
    "management_ip", "business_ip", "bmc_ip", "serial_number", "project",
]


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Demo"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row.get(key, "") for key in HEADERS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = max(len(str(cell.value or "")) for cell in column) + 3
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    for number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[number].height = 22
    book.save(path)


def main() -> None:
    assets = [
        dict(asset_id="DEMO-ASSET-001", device_type="SERVER", vendor="DEMO-VENDOR",
             model="MODEL-A", status="ACTIVE", location="DEMO-RACK-A",
             management_ip="192.0.2.11", business_ip="198.51.100.11",
             bmc_ip="203.0.113.11", serial_number="DEMO-SN-001", project="DEMO-PROJECT"),
        dict(asset_id="DEMO-ASSET-002", device_type="SERVER", vendor="DEMO-VENDOR",
             model="MODEL-B", status="ACTIVE", location="DEMO-RACK-B",
             management_ip="192.0.2.12", business_ip="198.51.100.12",
             bmc_ip="203.0.113.12", serial_number="DEMO-SN-002", project="DEMO-PROJECT"),
        dict(asset_id="DEMO-ASSET-003", device_type="SERVER", vendor="DEMO-VENDOR",
             model="MODEL-C", status="ACTIVE", location="DEMO-RACK-C",
             management_ip="192.0.2.13", business_ip="198.51.100.13",
             bmc_ip="203.0.113.13", serial_number="DEMO-SN-003", project="DEMO-PROJECT"),
    ]
    bm = [
        dict(asset_id=" demo-asset-001 ", serial_number="demo-sn-001",
             bmc_ip="203.0.113.11"),
        dict(asset_id="DEMO-ASSET-002", serial_number="DEMO-SN-002",
             bmc_ip="203.0.113.99"),
        dict(asset_id="DEMO-ASSET-003", serial_number="DEMO-SN-OTHER",
             bmc_ip="203.0.113.13"),
    ]
    collected = [
        dict(asset_id="DEMO-ASSET-001", serial_number="DEMO-SN-001",
             management_ip="192.0.2.11", bmc_ip="203.0.113.11", model="MODEL-A"),
        dict(asset_id="DEMO-ASSET-002", serial_number="DEMO-SN-002",
             management_ip="192.0.2.12", bmc_ip="203.0.113.12", model="MODEL-X"),
        dict(asset_id="DEMO-ASSET-004", serial_number="DEMO-SN-004",
             management_ip="192.0.2.14", bmc_ip="203.0.113.14", model="MODEL-D"),
        dict(asset_id="DEMO-ASSET-005", serial_number="DEMO-SN-005",
             management_ip="192.0.2.15", bmc_ip="203.0.113.15", model="MODEL-E"),
        dict(asset_id="DEMO-ASSET-003", serial_number="DEMO-SN-003",
             management_ip="192.0.2.13", bmc_ip="not-an-ip", model="MODEL-C"),
        dict(asset_id="DEMO-ASSET-004", serial_number="DEMO-SN-004",
             management_ip="192.0.2.14", bmc_ip="203.0.113.14", model="MODEL-D"),
    ]
    write_xlsx(HERE / "assets.xlsx", assets)
    write_xlsx(HERE / "bm.xlsx", bm)
    with (HERE / "collected.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(collected)
    print("已生成 demo/assets.xlsx、demo/bm.xlsx、demo/collected.csv")


if __name__ == "__main__":
    main()
