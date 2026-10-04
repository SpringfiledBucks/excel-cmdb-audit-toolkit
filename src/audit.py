"""Read rows, normalize identifiers, reconcile against assets and write reports."""

from __future__ import annotations

import csv
import ipaddress
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


FIELDS = (
    "asset_id", "device_type", "vendor", "model", "status", "location",
    "management_ip", "business_ip", "bmc_ip", "serial_number", "project",
)
IDENTIFIERS = ("serial_number", "asset_id", "management_ip", "bmc_ip")
IP_FIELDS = {"management_ip", "business_ip", "bmc_ip"}
UPPER_FIELDS = {"asset_id", "serial_number", "model"}
NULLS = {"", "null", "none", "nan", "n/a", "na", "-"}
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
REPORT_FIELDS = (
    "source", "source_row", "status", "matched_asset_row", "match_key",
    "asset_id", "serial_number", "management_ip", "bmc_ip",
    "field", "old_value", "new_value", "reason",
)


@dataclass
class Record:
    source: str
    source_row: int
    data: dict[str, str]
    invalid: list[str] = field(default_factory=list)
    duplicate: list[str] = field(default_factory=list)


def _cell(value: object) -> str:
    if value is None:
        return ""
    value = str(value).strip()
    return "" if value.casefold() in NULLS else value


def _normalize(field_name: str, value: object) -> tuple[str, bool]:
    value = _cell(value)
    if not value:
        return "", False
    if field_name in IP_FIELDS:
        try:
            return str(ipaddress.ip_address(value)), False
        except ValueError:
            return value, True
    if field_name in UPPER_FIELDS:
        return value.upper(), False
    if field_name in {"vendor", "device_type", "status"}:
        return value.casefold(), False
    return value, False


def _read_rows(path: Path, source: str) -> list[Record]:
    if not path.is_file():
        raise ValueError(f"找不到 {source} 文件：{path}")
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.reader(file)
            try:
                rows = list(reader)
            except csv.Error as exc:
                raise ValueError(f"{source} CSV 无法解析：{exc}") from exc
    elif suffix == ".xlsx":
        # Preserve formula text as data; cached formula values may be absent and
        # would otherwise turn an input cell into an unexplained blank.
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        raise ValueError(f"{source} 仅支持 .csv 或 .xlsx：{path}")
    if not rows:
        raise ValueError(f"{source} 缺少表头")
    headers = [_cell(value).casefold() for value in rows[0]]
    if not any(header in FIELDS for header in headers):
        raise ValueError(f"{source} 没有可识别字段；可用字段：{', '.join(FIELDS)}")
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"{source} 表头存在空列名或重复列名")
    records: list[Record] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if len(values) > len(headers) and any(_cell(v) for v in values[len(headers):]):
            raise ValueError(f"{source} 第 {row_number} 行有表头之外的非空数据")
        raw = dict(zip(headers, values))
        normalized: dict[str, str] = {}
        invalid: list[str] = []
        for name in FIELDS:
            value, bad = _normalize(name, raw.get(name))
            normalized[name] = value
            if bad:
                invalid.append(name)
        records.append(Record(source, row_number, normalized, invalid))
    return records


def _mark_duplicates(records: Iterable[Record]) -> None:
    grouped: dict[tuple[str, str], list[Record]] = defaultdict(list)
    for record in records:
        for identifier in IDENTIFIERS:
            value = record.data[identifier]
            if value and identifier not in record.invalid:
                grouped[(identifier, value)].append(record)
    for (identifier, _), group in grouped.items():
        if len(group) > 1:
            for record in group:
                record.duplicate.append(identifier)


def _detail(record: Record, status: str, *, asset: Record | None = None,
            key: str = "", field_name: str = "", old: str = "",
            new: str = "", reason: str = "") -> dict[str, object]:
    return {
        "source": record.source, "source_row": record.source_row,
        "status": status, "matched_asset_row": asset.source_row if asset else "",
        "match_key": key, "asset_id": record.data["asset_id"],
        "serial_number": record.data["serial_number"],
        "management_ip": record.data["management_ip"], "bmc_ip": record.data["bmc_ip"],
        "field": field_name, "old_value": old, "new_value": new,
        "reason": reason,
    }


def _local_issues(record: Record) -> list[dict[str, object]]:
    issues = [_detail(record, "INVALID", field_name=name,
                      new=record.data[name], reason="invalid_ip") for name in record.invalid]
    issues.extend(_detail(record, "DUPLICATE", field_name=name,
                          new=record.data[name], reason=f"duplicate_identifier_in_{record.source}")
                  for name in record.duplicate)
    return issues


def _audit_reference(record: Record, collected_rows: set[int]) -> list[dict[str, object]]:
    issues = _local_issues(record)
    if issues:
        return issues
    if not any(record.data[name] for name in IDENTIFIERS):
        return [_detail(record, "MISSING", reason="no_identifier_in_assets")]
    if record.source_row not in collected_rows:
        return [_detail(record, "MISSING", reason="not_observed_in_collected")]
    return [_detail(record, "MATCHED", reason="reference_observed_by_collected")]


def _audit_candidate(record: Record, index: dict[str, dict[str, list[Record]]]) -> list[dict[str, object]]:
    issues = _local_issues(record)
    if issues:
        return issues
    candidates: dict[str, list[Record]] = {}
    for name in IDENTIFIERS:
        value = record.data[name]
        if value:
            matches = index[name].get(value, [])
            if matches:
                candidates[name] = matches
    ambiguous = [(name, matches) for name, matches in candidates.items() if len(matches) > 1]
    if ambiguous:
        return [_detail(record, "DUPLICATE", field_name=name,
                        new=record.data[name], reason="identifier_ambiguous_in_assets")
                for name, _ in ambiguous]
    targets = {matches[0].source_row for matches in candidates.values()}
    if len(targets) > 1:
        return [_detail(record, "CONFLICT", field_name=name,
                        old=f"assets row {matches[0].source_row}",
                        new=record.data[name], reason="identifiers_point_to_different_assets")
                for name, matches in candidates.items()]
    if not targets:
        reason = "no_identifier" if not any(record.data[name] for name in IDENTIFIERS) else "no_asset_match"
        return [_detail(record, "MISSING", reason=reason)]
    key = next(name for name in IDENTIFIERS if name in candidates)
    asset = candidates[key][0]
    asset_issues = [_detail(record, "INVALID", asset=asset, key=key, field_name=name,
                            old=asset.data[name], new=record.data[name],
                            reason="matched_asset_has_invalid_ip") for name in asset.invalid]
    asset_issues.extend(_detail(record, "DUPLICATE", asset=asset, key=key, field_name=name,
                                old=asset.data[name], new=record.data[name],
                                reason="matched_asset_has_duplicate_identifier")
                        for name in asset.duplicate)
    if asset_issues:
        return asset_issues
    differences = []
    for name in FIELDS:
        new = record.data[name]
        old = asset.data[name]
        if new and new != old:
            differences.append(_detail(record, "CONFLICT", asset=asset, key=key,
                                       field_name=name, old=old, new=new,
                                       reason="value_differs_from_assets"))
    return differences or [_detail(record, "MATCHED", asset=asset, key=key,
                                   reason="identifiers_and_present_fields_agree")]


def _safe(value: object) -> object:
    """Keep spreadsheet viewers from interpreting user-controlled text as formulas."""
    if isinstance(value, str):
        value = ILLEGAL_XML.sub(lambda match: f"\\u{ord(match.group()):04x}", value)
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=REPORT_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({name: _safe(row[name]) for name in REPORT_FIELDS})


def _write_workbook(path: Path, sections: dict[str, list[dict[str, object]]],
                    summary: dict[str, object]) -> None:
    book = Workbook()
    first = book.active
    first.title = "Summary"
    first.append(["metric", "value"])
    first.append(["input_rows", summary["input_rows"]])
    for source, count in summary["source_rows"].items():
        first.append([f"source_rows.{source}", count])
    for status, count in summary["status_counts"].items():
        first.append([f"primary_status.{status}", count])
    for status, count in summary["affected_row_counts"].items():
        first.append([f"affected_rows.{status}", count])
    first.freeze_panes = "A2"
    first.auto_filter.ref = first.dimensions
    for title, rows in sections.items():
        sheet = book.create_sheet(title)
        sheet.append(REPORT_FIELDS)
        for row in rows:
            sheet.append([_safe(row[name]) for name in REPORT_FIELDS])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = (14, 13, 14, 19, 20, 22, 22, 21, 20, 20, 27, 27, 41)
        for number, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, number).column_letter].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for sheet in book:
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="243B53")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    first.column_dimensions["A"].width = 31
    first.column_dimensions["B"].width = 20
    book.save(path)


def run_audit(assets_path: Path, bm_path: Path, collected_path: Path,
              output_path: Path) -> dict[str, object]:
    """Audit every data row, then generate Excel, CSV and machine-readable summary."""
    sources = {
        "assets": _read_rows(Path(assets_path), "assets"),
        "bm": _read_rows(Path(bm_path), "bm"),
        "collected": _read_rows(Path(collected_path), "collected"),
    }
    for records in sources.values():
        _mark_duplicates(records)
    index: dict[str, dict[str, list[Record]]] = {name: defaultdict(list) for name in IDENTIFIERS}
    for asset in sources["assets"]:
        for name in IDENTIFIERS:
            value = asset.data[name]
            if value and name not in asset.invalid:
                index[name][value].append(asset)
    candidate_details: dict[str, dict[int, list[dict[str, object]]]] = {
        "bm": {}, "collected": {}
    }
    observed_collected: set[int] = set()
    for source in ("bm", "collected"):
        for record in sources[source]:
            rows = _audit_candidate(record, index)
            candidate_details[source][record.source_row] = rows
            if source == "collected" and rows[0]["status"] in {"MATCHED", "CONFLICT"}:
                linked = rows[0]["matched_asset_row"]
                if isinstance(linked, int):
                    observed_collected.add(linked)
    details: list[dict[str, object]] = []
    row_status: Counter[str] = Counter()
    per_source: dict[str, Counter[str]] = {source: Counter() for source in sources}
    for source, records in sources.items():
        source_details = (
            [_audit_reference(record, observed_collected) for record in records]
            if source == "assets" else None
        )
        for position, record in enumerate(records):
            if source_details is None:
                rows = candidate_details[source][record.source_row]
            else:
                rows = source_details[position]
            status = str(rows[0]["status"])
            row_status[status] += 1
            per_source[source][status] += 1
            details.extend(rows)
    ordered_statuses = ("MATCHED", "CONFLICT", "MISSING", "DUPLICATE", "INVALID")
    summary: dict[str, object] = {
        "input_rows": sum(map(len, sources.values())),
        "source_rows": {source: len(records) for source, records in sources.items()},
        "status_counts": {status: row_status[status] for status in ordered_statuses},
        "source_status_counts": {
            source: {status: counts[status] for status in ordered_statuses}
            for source, counts in per_source.items()
        },
        "detail_rows": len(details),
        "affected_row_counts": {
            status: len({(row["source"], row["source_row"])
                         for row in details if row["status"] == status})
            for status in ordered_statuses
        },
        "collected_coverage": {
            "assets_observed": len(observed_collected),
            "assets_unobserved": sum(
                1 for asset in sources["assets"]
                if not asset.invalid and not asset.duplicate
                and any(asset.data[name] for name in IDENTIFIERS)
                and asset.source_row not in observed_collected
            ),
        },
        "count_basis": "status_counts and source_status_counts count one primary status per input row (INVALID > DUPLICATE > CONFLICT > MISSING > MATCHED). affected_row_counts count distinct input rows per detail status and can overlap; detail_rows count all exported details.",
    }
    output = Path(output_path)
    output.mkdir(parents=True, exist_ok=True)
    sections = {
        "Matched": [row for row in details if row["status"] == "MATCHED"],
        "Conflicts": [row for row in details if row["status"] in {"CONFLICT", "INVALID"}],
        "Missing": [row for row in details if row["status"] == "MISSING"],
        "Duplicates": [row for row in details if row["status"] == "DUPLICATE"],
    }
    _write_workbook(output / "audit_result.xlsx", sections, summary)
    for name, section in (("matched", "Matched"), ("conflicts", "Conflicts"),
                          ("missing", "Missing"), ("duplicates", "Duplicates")):
        _write_csv(output / f"{name}.csv", sections[section])
    (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
                                          encoding="utf-8")
    return summary
