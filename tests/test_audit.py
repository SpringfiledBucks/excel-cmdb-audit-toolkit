import csv
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.audit import REPORT_FIELDS, _safe, run_audit


HEADERS = ["asset_id", "serial_number", "management_ip", "bmc_ip", "model"]


def write_rows(path: Path, rows: list[dict[str, str]]) -> None:
    if path.suffix == ".csv":
        with path.open("w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=HEADERS)
            writer.writeheader()
            writer.writerows(rows)
    else:
        book = Workbook()
        sheet = book.active
        sheet.append(HEADERS)
        for row in rows:
            sheet.append([row.get(name, "") for name in HEADERS])
        book.save(path)


def audit(tmp_path: Path, assets: list[dict[str, str]],
          bm: list[dict[str, str]], collected: list[dict[str, str]]):
    tmp_path.mkdir(parents=True, exist_ok=True)
    paths = [tmp_path / "assets.xlsx", tmp_path / "bm.xlsx", tmp_path / "collected.csv"]
    for path, rows in zip(paths, (assets, bm, collected)):
        write_rows(path, rows)
    output = tmp_path / "output"
    result = run_audit(*paths, output)
    reports = {}
    for name in ("matched", "conflicts", "missing", "duplicates"):
        with (output / f"{name}.csv").open(encoding="utf-8-sig", newline="") as file:
            reports[name] = list(csv.DictReader(file))
    return result, reports, output


BASE = dict(asset_id="DEMO-A1", serial_number="DEMO-S1",
            management_ip="192.0.2.1", bmc_ip="203.0.113.1", model="M1")


def test_exact_match_and_normalization(tmp_path):
    collected = dict(asset_id=" demo-a1 ", serial_number="demo-s1",
                     management_ip="192.0.2.1", bmc_ip="203.0.113.1", model="m1")
    summary, reports, output = audit(tmp_path, [BASE], [], [collected])
    assert summary["status_counts"]["MATCHED"] == 2
    assert summary["source_rows"] == {"assets": 1, "bm": 0, "collected": 1}
    assert len(reports["matched"]) == 2
    assert reports["matched"][1]["match_key"] == "serial_number"
    assert reports["matched"][1]["source_row"] == "2"
    assert (output / "audit_result.xlsx").is_file()


@pytest.mark.parametrize("changed,expected", [
    ({"serial_number": "DEMO-OTHER"}, "serial_number"),
    ({"management_ip": "192.0.2.99"}, "management_ip"),
    ({"model": "M2"}, "model"),
])
def test_field_conflicts(tmp_path, changed, expected):
    row = {**BASE, **changed}
    summary, reports, _ = audit(tmp_path, [BASE], [], [row])
    assert summary["status_counts"]["CONFLICT"] == 1
    assert any(item["field"] == expected and item["old_value"] == BASE[expected]
               and item["new_value"] == changed[expected] for item in reports["conflicts"])


def test_different_identifiers_target_different_assets(tmp_path):
    other = dict(asset_id="DEMO-A2", serial_number="DEMO-S2",
                 management_ip="192.0.2.2", bmc_ip="203.0.113.2", model="M2")
    row = dict(asset_id="DEMO-A2", serial_number="DEMO-S1")
    summary, reports, _ = audit(tmp_path, [BASE, other], [], [row])
    assert summary["source_status_counts"]["collected"]["CONFLICT"] == 1
    assert {item["reason"] for item in reports["conflicts"]} == {"identifiers_point_to_different_assets"}
    assert all(not item["matched_asset_row"] for item in reports["conflicts"])
    assert summary["collected_coverage"]["assets_unobserved"] == 2


def test_duplicates_in_source_and_reference(tmp_path):
    duplicate = {**BASE, "asset_id": "DEMO-A2", "serial_number": "DEMO-S1"}
    summary, reports, _ = audit(tmp_path, [BASE, duplicate], [], [BASE, BASE])
    assert summary["source_status_counts"]["assets"]["DUPLICATE"] == 2
    assert summary["source_status_counts"]["collected"]["DUPLICATE"] == 2
    assert len(reports["duplicates"]) >= 4
    assert all(item["source"] in {"assets", "collected"} for item in reports["duplicates"])


def test_candidate_does_not_match_invalid_or_duplicate_reference(tmp_path):
    invalid = {**BASE, "bmc_ip": "not-an-ip"}
    summary, reports, _ = audit(tmp_path / "bad", [invalid], [], [dict(serial_number="DEMO-S1")])
    assert summary["source_status_counts"]["collected"]["INVALID"] == 1
    assert any(item["reason"] == "matched_asset_has_invalid_ip" for item in reports["conflicts"])
    duplicate = {**BASE, "asset_id": "DEMO-A2"}
    summary, reports, _ = audit(tmp_path / "dup", [BASE, duplicate], [], [dict(asset_id="DEMO-A1")])
    assert summary["source_status_counts"]["collected"]["DUPLICATE"] == 1
    assert any(item["reason"] == "matched_asset_has_duplicate_identifier"
               for item in reports["duplicates"])


def test_missing_and_unobserved_assets(tmp_path):
    unknown = dict(asset_id="DEMO-NOT-FOUND")
    summary, reports, _ = audit(tmp_path, [BASE], [], [unknown])
    assert summary["status_counts"]["MISSING"] == 2
    assert {item["reason"] for item in reports["missing"]} == {
        "not_observed_in_collected", "no_asset_match"
    }
    assert summary["collected_coverage"]["assets_unobserved"] == 1


def test_invalid_ip_and_empty_identifier_row(tmp_path):
    invalid = {**BASE, "bmc_ip": "999.1.1.1"}
    summary, reports, _ = audit(tmp_path, [BASE], [], [invalid, {}])
    assert summary["source_status_counts"]["collected"]["INVALID"] == 1
    assert summary["source_status_counts"]["collected"]["MISSING"] == 1
    assert any(item["reason"] == "invalid_ip" for item in reports["conflicts"])
    assert any(item["reason"] == "no_identifier" for item in reports["missing"])


def test_invalid_and_duplicate_both_reported_with_one_primary_status(tmp_path):
    row1 = {**BASE, "bmc_ip": "not-an-ip"}
    row2 = {**BASE, "asset_id": "DEMO-A2"}
    summary, reports, _ = audit(tmp_path, [row1, row2], [], [])
    assert summary["source_status_counts"]["assets"]["INVALID"] == 1
    assert summary["source_status_counts"]["assets"]["DUPLICATE"] == 1
    assert any(item["source_row"] == "2" and item["field"] == "bmc_ip"
               for item in reports["conflicts"])
    assert {item["source_row"] for item in reports["duplicates"]} == {"2", "3"}
    assert summary["affected_row_counts"]["DUPLICATE"] == 2
    assert summary["status_counts"]["DUPLICATE"] == 1
    assert sum(summary["status_counts"].values()) == summary["input_rows"]
    assert sum(map(len, reports.values())) == summary["detail_rows"]


def test_empty_reference_and_row_conservation(tmp_path):
    summary, reports, output = audit(tmp_path, [{}], [{}], [{}])
    assert summary["input_rows"] == 3
    assert summary["status_counts"]["MISSING"] == 3
    assert sum(summary["status_counts"].values()) == summary["input_rows"]
    assert sum(map(len, reports.values())) == summary["detail_rows"]
    assert {item["source"] for item in reports["missing"]} == {"assets", "bm", "collected"}
    assert json.loads((output / "summary.json").read_text(encoding="utf-8")) == summary
    book = load_workbook(output / "audit_result.xlsx", read_only=True)
    assert book.sheetnames == ["Summary", "Matched", "Conflicts", "Missing", "Duplicates"]
    assert [cell.value for cell in next(book["Conflicts"].rows)] == list(REPORT_FIELDS)


def test_formula_like_values_are_escaped_in_exports(tmp_path):
    row = {**BASE, "model": "=HYPERLINK(\"x\")"}
    _, reports, output = audit(tmp_path, [BASE], [], [row])
    assert reports["conflicts"][0]["new_value"].startswith("'=")
    book = load_workbook(output / "audit_result.xlsx")
    cell = book["Conflicts"].cell(2, 12)
    assert cell.data_type == "s"
    assert cell.value.startswith("'=")


def test_control_characters_and_indented_formula_are_safe(tmp_path):
    row = {**BASE, "model": "  =SUM(1,1)\x01"}
    _, reports, output = audit(tmp_path, [BASE], [], [row])
    assert reports["conflicts"][0]["new_value"] == "'=SUM(1,1)\\u0001"
    assert _safe("  =SUM(1,1)") == "'  =SUM(1,1)"
    book = load_workbook(output / "audit_result.xlsx")
    assert book["Conflicts"].cell(2, 12).data_type == "s"


def test_unrecognized_headers_fail_without_dropping_rows(tmp_path):
    path = tmp_path / "assets.csv"
    path.write_text("unknown\nvalue\n", encoding="utf-8")
    with pytest.raises(ValueError, match="没有可识别字段"):
        run_audit(path, path, path, tmp_path / "output")
